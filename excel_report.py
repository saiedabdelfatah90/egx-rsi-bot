import json
import os
from datetime import datetime, timedelta
from zoneinfo import ZoneInfo

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

CAIRO = ZoneInfo("Africa/Cairo")
FRAME_ORDER = ["15 دقيقة", "ساعة", "يومي"]

# لازم تطابق إعدادات bot.py (لو غيّرتها هناك، غيّرها هنا)
TARGETS = {"15 دقيقة": (3, 6), "ساعة": (5, 10), "يومي": (5, 10)}
STOP_PCT = 7

REASONS = {
    "target": "هدف",
    "t1+t2": "هدفين",
    "t2": "هدف ثاني",
    "stop": "وقف خسارة",
    "t1+stop": "هدف أول ثم وقف",
}

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(bold=True, size=14, color="1F4E78")
GREEN = PatternFill("solid", fgColor="C6EFCE")
RED = PatternFill("solid", fgColor="FFC7CE")
TOTAL_FILL = PatternFill("solid", fgColor="DDEBF7")
_side = Side(style="thin", color="BFBFBF")
BORDER = Border(left=_side, right=_side, top=_side, bottom=_side)
CENTER = Alignment(horizontal="center", vertical="center")
PCT = '0.0"%"'
PCT_SIGNED = '+0.00"%";-0.00"%";0.00"%"'


def load_json(path, default):
    if os.path.exists(path):
        try:
            with open(path, encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            return default
    return default


def load_trades(path="trades.json"):
    data = load_json(path, [])
    if not isinstance(data, list):
        return []
    return [t for t in data if t.get("frame") in FRAME_ORDER]


def stats(items):
    wins = [t.get("pnl", 0) for t in items if t.get("pnl", 0) > 0]
    losses = [t.get("pnl", 0) for t in items if t.get("pnl", 0) <= 0]
    total = len(items)
    rate = (len(wins) / total * 100) if total else 0
    return len(wins), len(losses), rate, sum(wins), sum(losses)


def period_items(trades, period):
    now = datetime.now(CAIRO)
    today = now.strftime("%Y-%m-%d")
    if period == "today":
        items = [t for t in trades if t.get("date") == today]
        title = f"اليوم {today}"
    elif period == "week":
        start = (now - timedelta(days=6)).strftime("%Y-%m-%d")
        items = [t for t in trades if t.get("date", "") >= start]
        title = "آخر 7 أيام"
    else:
        items = list(trades)
        title = "كل الصفقات"
    return items, title


def _header(ws, row, headers):
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=col, value=h)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = CENTER
        c.border = BORDER


def _widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _summary_sheet(wb, items, title):
    ws = wb.active
    ws.title = "الملخص"
    ws.sheet_view.rightToLeft = True
    ws["A1"] = f"تقرير الصفقات — {title}"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = "تم الإنشاء: " + datetime.now(CAIRO).strftime("%Y-%m-%d %H:%M")

    headers = ["الفريم", "عدد الصفقات", "ناجحة", "نسبة النجاح",
               "مجموع الأرباح", "خاسرة", "مجموع الخسائر", "الصافي"]
    _header(ws, 4, headers)

    rows = [(fr, [t for t in items if t.get("frame") == fr]) for fr in FRAME_ORDER]
    rows.append(("الإجمالي", items))
    r = 5
    for name, its in rows:
        w, l, rate, sw, sl = stats(its)
        values = [name, len(its), w, rate, sw, l, sl, sw + sl]
        for col, v in enumerate(values, start=1):
            c = ws.cell(row=r, column=col, value=v)
            c.border = BORDER
            c.alignment = CENTER
            if name == "الإجمالي":
                c.fill = TOTAL_FILL
                c.font = Font(bold=True)
        ws.cell(row=r, column=4).number_format = PCT
        ws.cell(row=r, column=5).number_format = PCT_SIGNED
        ws.cell(row=r, column=7).number_format = PCT_SIGNED
        net = ws.cell(row=r, column=8)
        net.number_format = PCT_SIGNED
        if its:
            net.fill = GREEN if (sw + sl) > 0 else RED
        r += 1

    _widths(ws, [16, 14, 10, 14, 16, 10, 16, 12])
    ws.freeze_panes = "A5"


def _trades_sheet(wb, items):
    ws = wb.create_sheet("الصفقات المقفولة")
    ws.sheet_view.rightToLeft = True
    headers = ["التاريخ", "الوقت", "السهم", "الفريم",
               "سعر الشراء", "سعر البيع", "النتيجة", "سبب القفل"]
    _header(ws, 1, headers)

    ordered = sorted(items, key=lambda t: (t.get("date", ""), t.get("time", "")), reverse=True)
    for i, t in enumerate(ordered, start=2):
        pnl = t.get("pnl", 0)
        values = [t.get("date", ""), t.get("time", ""), t.get("symbol", ""),
                  t.get("frame", ""), t.get("entry", 0), t.get("exit", 0), pnl,
                  REASONS.get(t.get("reason", ""), t.get("reason", ""))]
        for col, v in enumerate(values, start=1):
            c = ws.cell(row=i, column=col, value=v)
            c.border = BORDER
            c.alignment = CENTER
        ws.cell(row=i, column=5).number_format = "0.00"
        ws.cell(row=i, column=6).number_format = "0.00"
        res = ws.cell(row=i, column=7)
        res.number_format = PCT_SIGNED
        res.fill = GREEN if pnl > 0 else RED

    _widths(ws, [13, 9, 11, 12, 13, 13, 12, 18])
    ws.freeze_panes = "A2"
    if ordered:
        ws.auto_filter.ref = f"A1:H{len(ordered) + 1}"


def _open_sheet(wb, state):
    ws = wb.create_sheet("الصفقات المفتوحة")
    ws.sheet_view.rightToLeft = True
    headers = ["السهم", "الفريم", "أجزاء الشراء", "سعر الشراء الأول",
               "متوسط الشراء", "تم بيع جزء؟", "الوقف", "الهدف الأول", "الهدف الثاني"]
    _header(ws, 1, headers)

    rows = []
    for key, pos in state.items():
        if not isinstance(pos, dict) or pos.get("status") not in ("HALF", "FULL") or "|" not in key:
            continue
        sym, fr = key.split("|", 1)
        if fr in TARGETS:
            rows.append((FRAME_ORDER.index(fr), sym, fr, pos))
    rows.sort(key=lambda x: (x[0], x[1]))

    for i, (_, sym, fr, pos) in enumerate(rows, start=2):
        t1p, t2p = TARGETS[fr]
        avg = pos["avg"]
        first = pos["first"]
        values = [sym, fr, 2 if pos.get("status") == "FULL" else 1, first, avg,
                  "نعم (نصف الكمية)" if pos.get("t1_done") else "لا",
                  first * (1 - STOP_PCT / 100),
                  avg * (1 + t1p / 100), avg * (1 + t2p / 100)]
        for col, v in enumerate(values, start=1):
            c = ws.cell(row=i, column=col, value=v)
            c.border = BORDER
            c.alignment = CENTER
        for col in (4, 5, 7, 8, 9):
            ws.cell(row=i, column=col).number_format = "0.00"

    _widths(ws, [11, 12, 14, 16, 14, 16, 11, 13, 13])
    ws.freeze_panes = "A2"


def make_report(period, trades, state):
    items, title = period_items(trades, period)
    wb = Workbook()
    _summary_sheet(wb, items, title)
    _trades_sheet(wb, items)
    _open_sheet(wb, state)
    name = f"report_{period}_{datetime.now(CAIRO).strftime('%Y-%m-%d')}.xlsx"
    wb.save(name)
    return name, title, len(items)


def send_document(path, chat_id, caption=""):
    try:
        with open(path, "rb") as f:
            requests.post(
                f"https://api.telegram.org/bot{os.environ['BOT_TOKEN']}/sendDocument",
                data={"chat_id": chat_id, "caption": caption},
                files={"document": (os.path.basename(path), f)},
                timeout=60)
    except Exception as e:
        print("خطأ في إرسال الملف:", e)
